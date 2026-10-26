import io

from docx import Document

from spec_auditor.docx_export import build_docx
from spec_auditor.models import AuditReport, DefectType, Evidence, Finding, Severity
from spec_auditor.report import render_markdown


def finding(severity=Severity.HIGH, files=("pricing.md",), quote="Orders over RM100 receive a discount"):
    return Finding(
        evidence=[Evidence(file=f, quote=quote) for f in files],
        section="Discount Tiers",
        issue="The RM100 boundary is covered twice.",
        defect_type=DefectType.LOGICAL_ERROR,
        severity=severity,
        suggested_fix="State 'RM100 or more'.",
    )


def test_every_section_is_present():
    report = AuditReport(findings=[finding()], reviewed_files=["pricing.md"])

    markdown = render_markdown(report)

    for heading in ["A. Findings by file", "B. Findings by type", "C. Detailed findings",
                    "D. Summary", "E. Coverage ledger"]:
        assert heading in markdown


def test_headline_count_matches_the_findings():
    report = AuditReport(findings=[finding(), finding(Severity.LOW)], reviewed_files=["pricing.md"])

    assert "**2 finding(s)**" in render_markdown(report)


def test_pipes_in_content_are_escaped():
    hostile = Finding(
        evidence=[Evidence(file="a.md", quote="status is active | inactive | pending")],
        section="States",
        issue="Pipe | characters here",
        defect_type=DefectType.AMBIGUITY,
        severity=Severity.LOW,
        suggested_fix="Use a list | not pipes",
    )
    report = AuditReport(findings=[hostile], reviewed_files=["a.md"])

    rows = [line for line in render_markdown(report).splitlines() if line.startswith("| 1 ")]

    assert len(rows) == 1
    assert rows[0].count("|") - rows[0].count("\\|") == 9  # 8 columns means 9 delimiters


def test_files_without_findings_appear_in_the_ledger():
    report = AuditReport(findings=[finding()], reviewed_files=["pricing.md", "accounts.md"])

    markdown = render_markdown(report)

    assert "| accounts.md | Reviewed, findings: 0 |" in markdown
    assert "| pricing.md | Reviewed, findings: 1 |" in markdown


def test_docx_contains_the_findings_table():
    report = AuditReport(findings=[finding()], reviewed_files=["pricing.md"])

    document = Document(io.BytesIO(build_docx(report)))

    assert len(document.tables) == 2
    assert document.tables[0].cell(1, 5).text == "Logical error"
    assert document.tables[1].cell(1, 1).text == "Reviewed, findings: 1"


