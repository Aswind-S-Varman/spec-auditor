import io

from docx import Document
from docx.shared import Pt

from .models import AuditReport, Severity

SEVERITIES = [severity.value for severity in Severity]


def build_docx(report: AuditReport) -> bytes:
    """Built from the report object, not by converting the Markdown."""
    document = Document()

    document.add_heading("Specification Audit Report", level=0)
    severities = report.counts_by_severity()
    document.add_paragraph(
        f"{report.total} finding(s) across {len(report.reviewed_files)} file(s) — "
        f"High: {severities['High']}, Medium: {severities['Medium']}, Low: {severities['Low']}"
    )

    document.add_heading("Detailed findings", level=1)
    if report.findings:
        headers = ["#", "Files", "Section", "Evidence", "Issue", "Type", "Severity", "Suggested fix"]
        table = document.add_table(rows=1, cols=len(headers))
        table.style = "Table Grid"

        for cell, header in zip(table.rows[0].cells, headers):
            run = cell.paragraphs[0].add_run(header)
            run.bold = True

        for number, finding in enumerate(report.findings, start=1):
            evidence = "\n".join(f'{item.file}: "{item.quote}"' for item in finding.evidence)
            values = [
                str(number),
                ", ".join(finding.files),
                finding.section,
                evidence,
                finding.issue,
                finding.defect_type.value,
                finding.severity.value,
                finding.suggested_fix,
            ]
            for cell, value in zip(table.add_row().cells, values):
                cell.text = value
    else:
        document.add_paragraph("No findings.")

    document.add_heading("Coverage ledger", level=1)
    ledger = document.add_table(rows=1, cols=2)
    ledger.style = "Table Grid"
    for cell, header in zip(ledger.rows[0].cells, ["File", "Status"]):
        cell.paragraphs[0].add_run(header).bold = True

    for file, count in report.coverage_ledger().items():
        row = ledger.add_row().cells
        row[0].text = file
        row[1].text = f"Reviewed, findings: {count}"

    for style_name in ("Normal",):
        document.styles[style_name].font.size = Pt(10)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


